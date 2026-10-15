#!/usr/bin/env python3
"""Make a spreadsheet of your Spotify playlists with Bandcamp links.

One tab per playlist with Artist, Track, Year and a Bandcamp link when the
track is found there. The result is an .xlsx file: upload it to Google Drive
and open it with Google Sheets, and the tabs carry over.

Input is either Exportify CSV exports (https://exportify.app, no setup) or
your Spotify account through the API (needs a Spotify developer app).

Usage:
    python bandcamp_sheet.py --csv exports/*.csv     # Exportify files
    python bandcamp_sheet.py --list                  # your own playlists (API)
    python bandcamp_sheet.py                         # all playlists you own (API)
    python bandcamp_sheet.py "Road Trip" "Chill"     # selected playlists (API)
"""

import argparse
import csv
import json
import os
import re
import sys
import time
import unicodedata
import urllib.request
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
CACHE_FILE = HERE / ".bandcamp-cache.json"
BANDCAMP_SEARCH = "https://bandcamp.com/api/bcsearch_public_api/1/autocomplete_elastic"
REQUEST_DELAY = 1.5  # seconds between Bandcamp searches; be polite


# --- Spotify input -----------------------------------------------------------

def spotify_client():
    import spotipy
    from spotipy.oauth2 import SpotifyOAuth
    for var in ("SPOTIPY_CLIENT_ID", "SPOTIPY_CLIENT_SECRET"):
        if not os.environ.get(var):
            sys.exit(f"Missing {var}. See README.md, or use --csv with Exportify files.")
    auth = SpotifyOAuth(
        redirect_uri=os.environ.get("SPOTIPY_REDIRECT_URI", "http://127.0.0.1:8888/callback"),
        scope="playlist-read-private playlist-read-collaborative",
        cache_path=str(HERE / ".spotify-cache"),
    )
    return spotipy.Spotify(auth_manager=auth)


def paginate(sp, page):
    while page:
        yield from page["items"]
        page = sp.next(page) if page.get("next") else None


def api_playlists(args):
    sp = spotify_client()
    me = sp.current_user()["id"]
    owned = [p for p in paginate(sp, sp.current_user_playlists(limit=50))
             if p and p["owner"]["id"] == me]
    if args.list:
        for p in owned:
            print(f"{p['tracks']['total']:>5}  {p['name']}")
        sys.exit(0)
    if args.playlists:
        wanted = {w.lower() for w in args.playlists}
        owned = [p for p in owned if p["name"].lower() in wanted or p["id"].lower() in wanted]
        if not owned:
            sys.exit("No matching playlists. Run with --list to see names.")

    out = {}
    for p in owned:
        print(f"Reading {p['name']}…")
        tracks = []
        for item in paginate(sp, sp.playlist_items(p["id"], limit=100, additional_types=("track",))):
            t = item.get("track")
            if not t or t.get("type") != "track" or t.get("is_local"):
                continue
            tracks.append({
                "artist": ", ".join(a["name"] for a in t["artists"]),
                "title": t["name"],
                "year": (t["album"].get("release_date") or "")[:4],
            })
        out[p["name"]] = tracks
    return out


def pick(row, *names):
    lower = {k.strip().lower(): v for k, v in row.items() if k}
    for n in names:
        if lower.get(n):
            return lower[n].strip()
    return ""


def csv_playlists(paths):
    out = {}
    for path in paths:
        tracks = []
        with open(path, newline="", encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                title = pick(row, "track name", "name", "title")
                artist = pick(row, "artist name(s)", "artist name", "artists", "artist")
                if title and artist:
                    tracks.append({
                        "artist": ", ".join(a.strip() for a in re.split(r";", artist)),
                        "title": title,
                        "year": pick(row, "release date", "album release date", "year")[:4],
                    })
        out[Path(path).stem.replace("_", " ")] = tracks
    return out


# --- Bandcamp lookup ---------------------------------------------------------

def normalize(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\s*[\(\[][^\)\]]*(remaster|version|edit|mix|feat|with)[^\)\]]*[\)\]]", "", s)
    s = re.sub(r"\s+-\s+.*(remaster|version|edit|mix).*$", "", s)
    s = re.sub(r"\b(feat|ft)\.?\s.*$", "", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def similarity(a, b):
    return SequenceMatcher(None, normalize(a), normalize(b)).ratio()


class Bandcamp:
    def __init__(self):
        self.cache = json.loads(CACHE_FILE.read_text()) if CACHE_FILE.exists() else {}
        self.last = 0.0
        self.failures = 0

    def search(self, text):
        wait = REQUEST_DELAY - (time.time() - self.last)
        if wait > 0:
            time.sleep(wait)
        self.last = time.time()
        body = json.dumps({"search_text": text, "search_filter": "t",
                           "full_page": False, "fan_id": None}).encode()
        req = urllib.request.Request(BANDCAMP_SEARCH, data=body, headers={
            "Content-Type": "application/json", "User-Agent": "Mozilla/5.0 (playlist-bandcamp-sheet)"})
        with urllib.request.urlopen(req, timeout=20) as r:
            return (json.load(r).get("auto") or {}).get("results") or []

    def find(self, track):
        """Return (status, url). status: found / possible / not found / error."""
        main_artist = track["artist"].split(",")[0].strip()
        key = f"{normalize(main_artist)}|{normalize(track['title'])}"
        if key in self.cache:
            return tuple(self.cache[key])
        try:
            results = self.search(f"{main_artist} {normalize(track['title'])}")
        except Exception as e:
            self.failures += 1
            print(f"    Bandcamp search failed: {e}")
            return "error", search_url(track)
        result = best_match(main_artist, track["title"], results)
        self.cache[key] = result
        CACHE_FILE.write_text(json.dumps(self.cache))
        return result


def best_match(artist, title, results):
    best, best_score = None, 0.0
    for r in results:
        if r.get("type") != "t":
            continue
        t_score = similarity(title, r.get("name", ""))
        a_score = max(similarity(artist, r.get("band_name", "")),
                      similarity(artist, r.get("artist_name", "") or r.get("artist", "")))
        if normalize(artist) and normalize(artist) in normalize(r.get("band_name", "")):
            a_score = 1.0
        score = 0.6 * t_score + 0.4 * a_score
        if t_score >= 0.8 and score > best_score:
            best, best_score = r, score
    if not best:
        return "not found", ""
    url = best.get("item_url_path") or best.get("url") or ""
    return ("found" if best_score >= 0.85 else "possible"), url


def search_url(track):
    return f"https://bandcamp.com/search?q={quote_plus(track['artist'].split(',')[0] + ' ' + track['title'])}&item_type=t"


# --- Spreadsheet -------------------------------------------------------------

def sheet_title(name, used):
    base = re.sub(r"[\[\]:*?/\\]", "-", name).strip()[:31] or "Playlist"
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title


def write_workbook(playlists, results, path):
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Playlist", "Tracks", "On Bandcamp", "Possible match", "Not found"])
    bold = Font(bold=True)
    link_font = Font(color="0563C1", underline="single")
    fills = {"possible": PatternFill("solid", fgColor="FFF2CC"),
             "not found": PatternFill("solid", fgColor="F2F2F2"),
             "error": PatternFill("solid", fgColor="FCE4E4")}
    used = {"summary"}

    for name, tracks in playlists.items():
        ws = wb.create_sheet(sheet_title(name, used))
        ws.append(["Artist", "Track", "Year", "Bandcamp", "Status"])
        counts = {"found": 0, "possible": 0, "not found": 0, "error": 0}
        for t in tracks:
            status, url = results[(t["artist"], t["title"])]
            counts[status] += 1
            ws.append([t["artist"], t["title"], int(t["year"]) if t["year"].isdigit() else t["year"],
                       url or "", status])
            row = ws.max_row
            if url:
                cell = ws.cell(row=row, column=4)
                cell.hyperlink, cell.font = url, link_font
            if status in fills:
                for col in range(1, 6):
                    ws.cell(row=row, column=col).fill = fills[status]
        for cell in ws[1]:
            cell.font = bold
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, width in zip("ABCDE", (30, 40, 8, 60, 11)):
            ws.column_dimensions[col].width = width
        summary.append([name, len(tracks), counts["found"], counts["possible"],
                        counts["not found"] + counts["error"]])

    for cell in summary[1]:
        cell.font = bold
    for i, width in enumerate((40, 10, 14, 16, 12), 1):
        summary.column_dimensions[get_column_letter(i)].width = width
    wb.save(path)


# --- Main --------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("playlists", nargs="*", help="playlist names (API mode; default: all you own)")
    ap.add_argument("--csv", nargs="+", metavar="FILE", help="Exportify CSV files instead of the Spotify API")
    ap.add_argument("--list", action="store_true", help="list playlists you own and exit (API mode)")
    ap.add_argument("--out", default="bandcamp-playlists.xlsx", help="output file (default: %(default)s)")
    args = ap.parse_args()

    playlists = csv_playlists(args.csv) if args.csv else api_playlists(args)
    unique = {(t["artist"], t["title"]): t for tracks in playlists.values() for t in tracks}
    print(f"{sum(map(len, playlists.values()))} tracks in {len(playlists)} playlists, "
          f"{len(unique)} unique — searching Bandcamp…")

    bc = Bandcamp()
    results = {}
    for i, (key, t) in enumerate(unique.items(), 1):
        results[key] = bc.find(t)
        mark = {"found": "✓", "possible": "?", "not found": "·", "error": "!"}[results[key][0]]
        print(f"  [{i}/{len(unique)}] {mark} {t['artist']} – {t['title']}")
        if bc.failures >= 5 and i == bc.failures:
            sys.exit("Bandcamp search isn't responding. Check your connection and try again.")

    write_workbook(playlists, results, args.out)
    print(f"\nWrote {args.out} — upload it to Google Drive and open with Google Sheets.")


if __name__ == "__main__":
    main()
